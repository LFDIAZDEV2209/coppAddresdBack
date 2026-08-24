"""Importador de datos de prueba de pacientes desde el Excel de prueba.

Fuente: US_Health_Database_200k.xlsx, hoja "Patient_Database_50k" (50 000 filas).

Uso (desde ai-service, donde hay venv/uv):
    uv run --with openpyxl --with psycopg[binary] python ../coppAddresdBack/scripts/import_patients_from_excel.py

Idempotente: los pacientes cuyo MRN ya existe en app.patient_profiles se omiten
(no se borra nada, no se duplica nada).

Normalización aplicada (el Excel NO es el contrato del modelo de dominio):
- Identidad: MRN -> medical_record_number (únicos; MRN duplicados del origen se omiten),
  nombres, fecha de nacimiento, sexo -> género (Female->Femenino, Male->Masculino).
- Demografía: ciudad, estado (código), zip, contacto de emergencia.
- Cobertura: aseguradora -> app.insurers (catálogo sembrado por migración) + member_id.
- Estilo de vida y antecedentes: smoking, alcohol, ejercicio, discapacidad,
  hospitalización, cirugías (columnas del paciente).
- Diagnósticos -> app.patient_diagnoses (ICD-10, is_primary) con FK al catálogo
  app.icd10_codes (get-or-create; la descripción vive en el catálogo).
- Medicamentos -> app.patient_medications con FK al catálogo app.medications
  (get-or-create; NDC/RxNorm/clase viven en el catálogo; se omite "No medication").
- Alergias -> app.patient_allergies con FK al catálogo app.allergens
  (get-or-create; se omite "No known allergies").
- Signos vitales -> app.vital_signs con unidades SI (cm, kg, °C); la fecha de
  medición usa la última visita cuando está disponible.
- Derivables NO persistidas: edad, IMC, nombre completo, descripción del estado.
- Campos de citas/proveedor (última visita, próxima cita, médico, NPI, conteo de
  visitas) NO se modelan: pertenecen al módulo de citas, no al de pacientes.
- SSN (enmascarado) no se importa: dato sensible sin uso clínico.
"""

from __future__ import annotations

import argparse
import sys
import uuid
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
import psycopg

EXCEL_PATH = r"C:\Users\Luis\Documents\CoppAddresd\Repos\US_Health_Database_200k.xlsx"
SHEET_NAME = "Patient_Database_50k"
DB_DSN = "host=localhost port=5432 dbname=coppaddresd user=app_user password=CoppAddresdDev!2026"

SEX_TO_GENDER = {"Female": "Femenino", "Male": "Masculino"}


def parse_date(value: str | None) -> str | None:
    """Convierte MM/DD/YYYY a ISO YYYY-MM-DD."""
    if value is None:
        return None
    try:
        return datetime.strptime(value.strip(), "%m/%d/%Y").date().isoformat()
    except (ValueError, AttributeError):
        return None


def to_decimal(value, factor: Decimal, precision: int = 1) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return (Decimal(str(value)) * factor).quantize(
            Decimal(10) ** -precision, rounding=ROUND_HALF_UP
        )
    except (ValueError, ArithmeticError):
        return None


def fahrenheit_to_celsius(value) -> Decimal | None:
    """Convierte °F a °C: (F - 32) * 5/9, redondeado a 1 decimal."""
    if value is None or value == "":
        return None
    try:
        return ((Decimal(str(value)) - Decimal(32)) * Decimal("0.5555555556")).quantize(
            Decimal("0.1"), rounding=ROUND_HALF_UP
        )
    except (ValueError, ArithmeticError):
        return None


def to_int(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(round(float(value)))
    except (ValueError, TypeError):
        return None


def clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def normalize_rxnorm(value) -> str | None:
    """Normaliza RxNorm: quita el prefijo 'RxNorm:', espacios y comas (typos
    del origen, ej. 'RxNorm:2200,644' -> '2200644')."""
    if value is None:
        return None
    text = value.replace("RxNorm:", "").replace(",", "").replace(" ", "")
    return text or None


def find_col(header: list[object], prefix: str) -> int:
    """Localiza una columna por prefijo (robusto ante caracteres corruptos del encabezado)."""
    for idx, name in enumerate(header):
        if name is not None and str(name).strip().startswith(prefix):
            return idx
    raise KeyError(f"No se encontró la columna con prefijo '{prefix}'. Header: {header}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", default=EXCEL_PATH)
    parser.add_argument("--dsn", default=DB_DSN)
    parser.add_argument("--limit", type=int, default=None, help="Importar solo las primeras N filas (debug).")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.excel, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"Hoja '{SHEET_NAME}' no encontrada. Hojas: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    rows = ws.iter_rows(values_only=True)
    header = list(next(rows))
    col = {name: idx for idx, name in enumerate(header)}
    temp_col = find_col(header, "Temp")

    conn = psycopg.connect(args.dsn)
    conn.autocommit = False

    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id, name FROM app.insurers")
            insurer_ids = {row_name: str(row_id) for row_id, row_name in cur.fetchall()}
            cur.execute("SELECT id, name FROM app.allergens")
            allergen_ids = {row_name: str(row_id) for row_id, row_name in cur.fetchall()}
            cur.execute("SELECT id, code FROM app.icd10_codes")
            icd10_ids = {row_code: str(row_id) for row_id, row_code in cur.fetchall()}
            cur.execute("SELECT id, name FROM app.medications")
            medication_ids = {row_name: str(row_id) for row_id, row_name in cur.fetchall()}
            cur.execute("SELECT medical_record_number FROM app.patient_profiles")
            existing_mrns = {row[0] for row in cur.fetchall()}

        if not insurer_ids:
            sys.exit("El catálogo app.insurers está vacío: aplica la migración AddPatientsModule.")

        patients: list[tuple] = []
        diagnoses: list[tuple] = []
        medications: list[tuple] = []
        allergies: list[tuple] = []
        vitals: list[tuple] = []
        new_allergens: set[str] = set()
        new_icd10: dict[str, str] = {}
        new_medications: dict[str, tuple] = {}
        skipped_dupes = 0
        row_index = 1

        for row_index, row in enumerate(rows, start=2):
            if args.limit and len(patients) >= args.limit:
                break

            mrn = clean(row[col["MRN"]])
            if mrn is None or mrn in existing_mrns:
                skipped_dupes += 1
                continue

            patient_id = str(uuid.uuid4())
            insurer_name = clean(row[col["Insurance Provider"]])
            insurer_id = insurer_ids.get(insurer_name) if insurer_name else None

            patients.append((
                patient_id,
                mrn,
                clean(row[col["First Name"]]),
                clean(row[col["Middle Name"]]),
                clean(row[col["Last Name"]]),
                parse_date(row[col["Date of Birth"]]),
                SEX_TO_GENDER.get(clean(row[col["Sex"]]), clean(row[col["Sex"]])),
                clean(row[col["Ethnicity"]]),
                clean(row[col["Blood Type"]]),
                None,  # teléfono: el origen no lo trae
                None,  # email: el origen no lo trae
                None,  # dirección: el origen no trae calle
                clean(row[col["City"]]),
                clean(row[col["State Code"]]),
                str(row[col["ZIP Code"]]).strip() if row[col["ZIP Code"]] is not None else None,
                clean(row[col["Emerg. Contact"]]),
                insurer_id,
                clean(row[col["Member ID"]]),
                clean(row[col["Smoking"]]),
                clean(row[col["Alcohol"]]),
                clean(row[col["Exercise"]]),
                clean(row[col["Disability"]]),
                clean(row[col["Hospitalization Hx"]]),
                clean(row[col["Surgery History"]]),
                "Activo",
            ))

            # Diagnósticos (el código se resuelve contra el catálogo al final;
            # la descripción solo alimenta el catálogo si el código es nuevo).
            primary_code = clean(row[col["ICD-10 Primary"]])
            if primary_code:
                if primary_code not in icd10_ids and primary_code not in new_icd10:
                    new_icd10[primary_code] = clean(row[col["Primary Diagnosis"]])
                diagnoses.append((str(uuid.uuid4()), patient_id, primary_code, True))
            secondary_code = clean(row[col["ICD-10 Secondary"]])
            if secondary_code:
                if secondary_code not in icd10_ids and secondary_code not in new_icd10:
                    new_icd10[secondary_code] = clean(row[col["Secondary Diagnosis"]])
                diagnoses.append((str(uuid.uuid4()), patient_id, secondary_code, False))

            # Medicamentos (1 y 2): NDC/RxNorm/clase alimentan el catálogo
            # solo si el nombre es nuevo; la frecuencia es por paciente.
            for slot in (1, 2):
                name = clean(row[col[f"Medication {slot}"]])
                if not name or name.lower() == "no medication":
                    continue
                ndc_v = clean(row[col[f"NDC Med {slot}"]])
                rxnorm_v = normalize_rxnorm(clean(row[col[f"RxNorm Med {slot}"]]))
                drugclass_v = clean(row[col[f"Drug Class Med {slot}"]]) if f"Drug Class Med {slot}" in col else None
                if name not in medication_ids:
                    entry = new_medications.setdefault(name, (None, None, None))
                    new_medications[name] = (entry[0] or ndc_v, entry[1] or rxnorm_v, entry[2] or drugclass_v)
                medications.append((
                    str(uuid.uuid4()), patient_id, name,
                    clean(row[col[f"Frequency Med {slot}"]]) if f"Frequency Med {slot}" in col else None,
                    slot,
                ))

            # Alergias (se omite "No known allergies")
            allergen = clean(row[col["Allergies"]])
            if allergen and allergen.lower() != "no known allergies":
                if allergen not in allergen_ids:
                    new_allergens.add(allergen)
                allergies.append((str(uuid.uuid4()), patient_id, allergen, None))

            # Signos vitales (unidades SI; medición asociada a la última visita)
            measured_at = parse_date(row[col["Last Visit"]]) or datetime.utcnow().date().isoformat()
            vitals.append((
                str(uuid.uuid4()), patient_id, measured_at,
                to_int(row[col["Systolic BP"]]),
                to_int(row[col["Diastolic BP"]]),
                to_int(row[col["Heart Rate"]]),
                fahrenheit_to_celsius(row[temp_col]),
                to_int(row[col["O2 Sat %"]]),
                to_decimal(row[col["Height (in)"]], Decimal("2.54")),
                to_decimal(row[col["Weight (lbs)"]], Decimal("0.45359237")),
            ))

            existing_mrns.add(mrn)

        print(f"Filas leídas: {row_index - 1}; pacientes a insertar: {len(patients)}; "
              f"omitidos (MRN duplicado ya existente): {skipped_dupes}")

        with conn.cursor() as cur:
            # Upsert de catálogos de referencia (3NF): solo entradas nuevas,
            # dentro de la misma transacción que los pacientes.
            if new_allergens:
                with cur.copy("COPY app.allergens (id, name) FROM STDIN") as copy:
                    for name in sorted(new_allergens):
                        copy.write_row((str(uuid.uuid4()), name))
            if new_icd10:
                with cur.copy("COPY app.icd10_codes (id, code, description) FROM STDIN") as copy:
                    for code, desc in new_icd10.items():
                        copy.write_row((str(uuid.uuid4()), code, desc))
            if new_medications:
                with cur.copy(
                    "COPY app.medications (id, name, ndc, rx_norm, drug_class) FROM STDIN"
                ) as copy:
                    for name, (ndc, rxnorm, drugclass) in new_medications.items():
                        copy.write_row((str(uuid.uuid4()), name, ndc, rxnorm, drugclass))

            cur.execute("SELECT id, name FROM app.allergens")
            allergen_ids = {row_name: str(row_id) for row_id, row_name in cur.fetchall()}
            cur.execute("SELECT id, code FROM app.icd10_codes")
            icd10_ids = {row_code: str(row_id) for row_id, row_code in cur.fetchall()}
            cur.execute("SELECT id, name FROM app.medications")
            medication_ids = {row_name: str(row_id) for row_id, row_name in cur.fetchall()}

            with cur.copy(
                "COPY app.patient_profiles (id, medical_record_number, first_name, middle_name, "
                "last_name, date_of_birth, gender, ethnicity, blood_type, phone, email, address, "
                "city, state, postal_code, emergency_contact, insurer_id, member_id, "
                "smoking_status, alcohol_status, exercise_level, disability, "
                "hospitalization_history, surgery_history, status) FROM STDIN"
            ) as copy:
                for patient in patients:
                    copy.write_row(patient)

            with cur.copy(
                "COPY app.patient_diagnoses (id, patient_id, icd10_code_id, is_primary) FROM STDIN"
            ) as copy:
                for d_id, p_id, code, is_primary in diagnoses:
                    copy.write_row((d_id, p_id, icd10_ids[code], is_primary))

            with cur.copy(
                "COPY app.patient_medications (id, patient_id, medication_id, frequency, sort_order) "
                "FROM STDIN"
            ) as copy:
                for m_id, p_id, name, frequency, sort_order in medications:
                    copy.write_row((m_id, p_id, medication_ids[name], frequency, sort_order))

            with cur.copy(
                "COPY app.patient_allergies (id, patient_id, allergen_id, notes) FROM STDIN"
            ) as copy:
                for a_id, p_id, allergen, notes in allergies:
                    copy.write_row((a_id, p_id, allergen_ids[allergen], notes))

            with cur.copy(
                "COPY app.vital_signs (id, patient_id, measured_at, systolic, diastolic, "
                "heart_rate, temperature_c, o2_saturation, height_cm, weight_kg) FROM STDIN"
            ) as copy:
                for v in vitals:
                    copy.write_row(v)

            conn.commit()

            cur.execute("SELECT count(*) FROM app.patient_profiles")
            total_patients = cur.fetchone()[0]
            cur.execute("SELECT count(*) FROM app.patient_diagnoses")
            total_diagnoses = cur.fetchone()[0]
            cur.execute("SELECT count(*) FROM app.patient_medications")
            total_medications = cur.fetchone()[0]
            cur.execute("SELECT count(*) FROM app.patient_allergies")
            total_allergies = cur.fetchone()[0]
            cur.execute("SELECT count(*) FROM app.vital_signs")
            total_vitals = cur.fetchone()[0]
            cur.execute("SELECT count(*) FROM app.allergens")
            total_allergens = cur.fetchone()[0]
            cur.execute("SELECT count(*) FROM app.icd10_codes")
            total_icd10 = cur.fetchone()[0]
            cur.execute("SELECT count(*) FROM app.medications")
            total_meds = cur.fetchone()[0]

        print("Totales en BD:")
        print(f"  patient_profiles: {total_patients}")
        print(f"  patient_diagnoses: {total_diagnoses}")
        print(f"  patient_medications: {total_medications}")
        print(f"  patient_allergies: {total_allergies}")
        print(f"  vital_signs: {total_vitals}")
        print(f"  catálogo allergens: {total_allergens}")
        print(f"  catálogo icd10_codes: {total_icd10}")
        print(f"  catálogo medications: {total_meds}")
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


if __name__ == "__main__":
    main()